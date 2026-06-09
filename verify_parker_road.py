"""Parker Road Plaza와 핵심 투자 로직 교차검증.

수업자료 Q16_2_Parker Road Plaza.xlsx를 기준으로 PMT, 대출잔액,
연도별 이자, IRR 계산을 검증한다. 추가로 이 프로젝트의 확정 규칙
(보증금 처리, LTR 운영비 12%, 손익분기 판정)을 회귀 테스트한다.
"""

from pathlib import Path

import numpy_financial as npf
import openpyxl

from calculations import calculate_dcf_irr, calculate_ltr_noi
from judge_property import judge_property
from mortgage import annual_debt_service, cpm_schedule, loan_balance_at


ROOT = Path(__file__).resolve().parent
PARKER_FILE = ROOT / "reference" / "Q16_2_Parker Road Plaza.xlsx"


def assert_close(name: str, actual: float, expected: float, tol: float = 1e-6) -> None:
    """허용오차 이내인지 확인하고 사람이 읽기 쉬운 결과를 출력한다."""
    diff = actual - expected
    status = "OK" if abs(diff) <= tol else "FAIL"
    print(f"[{status}] {name}: actual={actual:.12f}, expected={expected:.12f}, diff={diff:.12g}")
    if status == "FAIL":
        raise AssertionError(f"{name} mismatch: actual={actual}, expected={expected}, diff={diff}")


def verify_parker_workbook() -> None:
    """Parker Road Plaza workbook의 Excel 공식 결과와 Python 계산을 비교한다."""
    wb = openpyxl.load_workbook(PARKER_FILE, data_only=True)
    ws = wb["(a) (b) (c)"]

    principal = float(ws["F34"].value)
    annual_rate = float(ws["F12"].value)
    loan_years = int(ws["F10"].value)

    print("\n=== Parker Road Plaza: loan engine ===")
    assert_close(
        "annual debt service",
        annual_debt_service(principal, annual_rate, loan_years),
        float(ws["C64"].value),
    )

    schedule = cpm_schedule(principal, annual_rate, loan_years, round_output=False)
    for col, after_year in [("C", 1), ("D", 2), ("E", 3), ("F", 4), ("G", 5)]:
        assert_close(
            f"loan balance after year {after_year}",
            loan_balance_at(principal, annual_rate, loan_years, after_year),
            float(ws[f"{col}65"].value),
        )

        start = (after_year - 1) * 12 + 1
        end = after_year * 12
        py_interest = schedule.loc[
            (schedule["회차"] >= start) & (schedule["회차"] <= end),
            "이자",
        ].sum()
        assert_close(
            f"annual interest year {after_year}",
            py_interest,
            float(ws[f"{col}66"].value),
        )

    print("\n=== Parker Road Plaza: IRR engine ===")
    btcf = [float(ws.cell(143, col).value) for col in range(3, 10)]
    atcf = [float(ws.cell(144, col).value) for col in range(3, 10)]
    assert_close("BTCF IRR", npf.irr(btcf), float(ws["C146"].value))
    assert_close("ATCF IRR", npf.irr(atcf), float(ws["C147"].value))


def verify_project_rules() -> None:
    """확정된 프로젝트 로직 규칙을 작은 예제로 회귀 검증한다."""
    params = {
        "purchase_price": 300_000_000,
        "ltv": 0.60,
        "hold_years": 5,
        "tax_rate": 0.35,
        "cap_rate": 0.05,
        "building_ratio": 0.80,
    }
    mortgage = {
        "principal": 180_000_000,
        "annual_rate": 0.045,
        "loan_years": 20,
    }

    print("\n=== Project rules regression ===")
    base = calculate_dcf_irr(10_000_000, params, mortgage)
    with_deposit = calculate_dcf_irr(10_000_000, params, mortgage, deposit=50_000_000)
    with_setup = calculate_dcf_irr(
        10_000_000,
        {**params, "setup_cost": 7_000_000},
        mortgage,
    )

    assert_close("deposit reduces equity 1:1", base["equity"] - with_deposit["equity"], 50_000_000)
    assert_close("setup cost increases equity 1:1", with_setup["equity"] - base["equity"], 7_000_000)

    ltr_noi = calculate_ltr_noi(monthly_rent=1_000_000, vacancy_rate=0.05, opex_ratio=0.12)
    assert_close("LTR NOI uses 12% opex", ltr_noi["noi"], 12_000_000 * 0.95 * 0.88)

    cases = [
        ("breakeven miss rejects", judge_property(20, 0.70, 0.60, 10_000_000, 5_000_000)["verdict"], "비추천"),
        ("IRR >= 15 recommends", judge_property(15, 0.50, 0.60, 10_000_000, 5_000_000)["verdict"], "추천"),
        ("IRR 5-15 cautions", judge_property(7, 0.50, 0.60, 10_000_000, 5_000_000)["verdict"], "신중"),
        ("IRR < 5 rejects", judge_property(4.9, 0.50, 0.60, 10_000_000, 5_000_000)["verdict"], "비추천"),
    ]
    for name, actual, expected in cases:
        status = "OK" if actual == expected else "FAIL"
        print(f"[{status}] {name}: actual={actual}, expected={expected}")
        if status == "FAIL":
            raise AssertionError(f"{name} mismatch: actual={actual}, expected={expected}")


def main() -> None:
    verify_parker_workbook()
    verify_project_rules()
    print("\nAll Parker Road and project rule checks passed.")


if __name__ == "__main__":
    main()
